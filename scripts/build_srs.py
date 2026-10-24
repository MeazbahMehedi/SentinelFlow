"""Rebuild the Word specification from the delivered design and requirements.
Optional documentation dependency: python-docx. Does not affect app startup.
"""
from pathlib import Path
from datetime import date
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT=Path(__file__).resolve().parents[1]
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=Inches(.75);sec.bottom_margin=Inches(.7)
sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
sec.footer_distance=Inches(.3)
styles=doc.styles
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    styles[name].font.name='Calibri'
    styles[name].font.color.rgb=RGBColor(0,0,0)
    styles[name].paragraph_format.space_after=Pt(7)
# Remove template paragraph borders, including Word's blue Title rule.
for style in styles:
    props=style.element.find(qn('w:pPr'))
    if props is not None:
        for border in list(props.findall(qn('w:pBdr'))):
            props.remove(border)
styles['Normal'].font.size=Pt(10.5)
styles['Normal'].paragraph_format.line_spacing=1.12
styles['Title'].font.size=Pt(29)
styles['Title'].font.bold=True
styles['Heading 1'].font.size=Pt(19)
styles['Heading 1'].paragraph_format.space_after=Pt(15)
styles['Heading 2'].font.size=Pt(12)
styles['Heading 2'].paragraph_format.space_before=Pt(10)
styles['Heading 2'].paragraph_format.space_after=Pt(6)
styles['Caption'].font.size=Pt(9)
styles['Caption'].font.italic=True
styles['Caption'].paragraph_format.space_after=Pt(10)
styles['Subtitle'].font.size=Pt(13)
foot=sec.footer.paragraphs[0]
foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=foot.add_run('SentinelFlow SRS  |  ');r.font.size=Pt(8)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');foot._p.append(f)
doc.core_properties.title='SentinelFlow Software Requirements Specification'
doc.core_properties.subject='Real Time Network Threat Detection with Live Dashboard and PCAP Replay'
doc.core_properties.author='SentinelFlow project team'
doc.core_properties.keywords='SRS network security PCAP replay Python Flask'


def p(text,style=None):return doc.add_paragraph(text,style)
def h(text):return doc.add_heading(text,2)
def page(title):doc.add_page_break();doc.add_heading(title,1)
def req(identifier,title,text):
    paragraph=doc.add_paragraph()
    paragraph.add_run(identifier+'  '+title+'  ').bold=True
    paragraph.add_run(text)
    return paragraph

def table(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    if widths:
        for c,w in zip(t.columns,widths):c.width=Inches(w)
    for c,title in zip(t.rows[0].cells,headers):c.text=str(title)
    for row in rows:
        cells=t.add_row().cells
        for c,value in zip(cells,row):c.text=str(value)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr()
        cant=OxmlElement('w:cantSplit');trpr.append(cant)
        if ri==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for ci,c in enumerate(row.cells):
            if widths:c.width=Inches(widths[ci])
            tcpr=c._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'80');e.set(qn('w:type'),'dxa');margins.append(e)
            tcpr.append(margins)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            tcpr.append(borders)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'263A35' if ri==0 else ('F2F6F4' if ri%2==0 else 'FFFFFF'));tcpr.append(shade)
            for par in c.paragraphs:
                par.paragraph_format.space_after=Pt(2);par.paragraph_format.line_spacing=1.05
                for run in par.runs:
                    run.font.size=Pt(9)
                    if ri==0:run.font.bold=True;run.font.color.rgb=RGBColor(255,255,255)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return t

def figure(filename,caption,max_h=4.6):
    path=ROOT/'docs'/'diagrams'/filename
    im=Image.open(path);w,hh=im.size
    width=min(6.55,max_h*w/hh)
    par=doc.add_paragraph();par.alignment=WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_after=Pt(4)
    run=par.add_run();run.add_picture(str(path),width=Inches(width))
    par.paragraph_format.keep_with_next=True
    cap=p(caption,'Caption');cap.alignment=WD_ALIGN_PARAGRAPH.CENTER

# Cover
p('SOFTWARE REQUIREMENTS SPECIFICATION').runs[0].font.size=Pt(10)
p('SentinelFlow Software Requirements Specification','Title')
p('Real Time Network Threat Detection System with Live Dashboard and PCAP Replay','Subtitle')
p('Release 1.1.0\n24 October 2026')
doc.add_paragraph()
h('Purpose of this specification')
p('This specification defines the delivered SentinelFlow application, its observable behavior, data structures, interfaces, security controls, and acceptance criteria. It supports implementation review, final project evaluation, maintenance, and an evidence based demonstration of the system.')
p('SentinelFlow provides a local Python workspace for replaying packet captures, monitoring visible network traffic, detecting suspicious behavior, and investigating alerts. Its six configurable detectors produce evidence that a user can inspect; their alerts are indicators requiring context rather than proof of compromise.')
h('Project authors')
table(['Field','Developer one','Developer two'],[['Full name','To be completed','To be completed'],['Student ID','To be completed','To be completed'],['Role and contribution','To be completed','To be completed'],['Institution and course','To be completed','To be completed']],[1.55,2.5,2.5])
p('Author identities can also be edited inside the application through The Developers page. This Word document remains separately editable for submission details.')

page('Document control and contents')
table(['Item','Value'],[['Document ID','SF SRS 001'],['Version','1.1.0'],['Baseline','Delivered source code and release verification']],[1.45,5.1])
h('Contents')
contents=[('1','Introduction and scope'),('2','Product context and actors'),('3','Use cases'),('4','Software architecture'),('5','Data flow and deployment'),('6','Account and workspace requirements'),('7','Capture ingestion requirements'),('8','Replay and session lifecycle'),('9','Behavioral detection requirements'),('10','DNS volume and integration requirements'),('11','Alert and host investigations'),('12','User interface requirements'),('13','Developer profiles and administration'),('14','Database design and data dictionary'),('15','External and internal interfaces'),('16','Security and quality requirements'),('17','Operational constraints and failures'),('18','Traceability and verification'),('19','Release validation and demonstration'),('20','Installation and acceptance'),('21','Limitations and future extensions'),('22','Accounts appearance and upgrade requirements'),('23','GitHub capture and shutdown requirements')]
table(['Section','Subject'],contents,[.65,5.9])

page('1 Introduction and scope')
h('Purpose and objectives')
p('Network investigations often require both live observations and repeatable analysis of recorded traffic. SentinelFlow combines those workflows in one local workspace. A user can move from a traffic overview to a specific detection, review the measurements that triggered it, associate it with an address, and record a decision.')
p('The project demonstrates Python backend engineering, event processing, relational persistence, input validation, interactive web interfaces, automated verification, and software documentation. A repeatable demonstration is a primary objective because classroom presentations should not depend on a live attack or unrestricted network access.')
h('Included scope')
p('The delivered system includes authenticated local access, ten separate navigation pages, PCAP import and timed replay, optional live capture, IPv4 and IPv6 metadata, six behavioral detectors, Suricata EVE alert import, persistent investigations, configurable rules and network ranges, developer profiles, private photo and CV uploads, exports, and an activity log.')
h('Excluded scope')
p('The baseline does not provide packet blocking, intrusion prevention, decryption, machine learning, distributed sensors, full TCP stream reconstruction, automatic remediation, email delivery, Zeek ingestion, or a production internet hosting configuration. A Suricata import reads existing alerts; it does not run that engine.')
h('Terminology')
table(['Term','Meaning'],[['PCAP','A file format containing recorded network packet data. PCAPNG is a newer block based capture format.'],['Event time','The timestamp attached to the original captured packet.'],['Replay speed','The factor used to scale scheduling delays; it does not change event timestamps.'],['Detector window','A capture time interval used to count rule observations.'],['Triage','Reviewing a signal and recording its investigation status.'],['Indicator','A configured address or observed behavior that merits review.'],['SSE','Server Sent Events, an HTTP stream used for browser updates.']],[1.25,5.3])

page('2 Product context and actors')
p('Registered users can import their own captures, replay traffic, configure their own rules, investigate alerts, and export their own data. The first account is an administrator. Additional accounts begin as analysts. Administrators can also capture host traffic, edit shared developer profiles, and change another account’s role.')
figure('01-context.png','Figure 1 System context and external inputs',max_h=3.0)
h('External sources and outputs')
p('The network interface supplies only traffic visible to the machine and requires an administrator account in the application. A PCAP file supplies historical observations without requiring capture privileges. A Suricata EVE file supplies signature alerts. Developer photos and CVs are user provided files. Outputs are browser views, persisted notes, CSV records, and JSON summaries.')
h('Assumptions and dependencies')
p('The user controls or is authorized to use the local service and has authorization to inspect the chosen traffic. Python 3.11 or newer and a modern browser are available. Initial dependency installation requires internet access. Dashboard assets are local; the optional GitHub repository view requires internet access and no API key. Physical live capture depends on OS permissions, sensor placement, and Npcap on Windows.')
h('Stakeholder outcomes')
p('An evaluator should be able to reproduce the supplied demonstration, explain why a detector fired, verify that replay speed preserves results, inspect the source layout, and identify the difference between implemented behavior and future extensions. The project team should be able to add its names, contributions, and CVs without editing templates.')

page('3 Use cases')
table(['ID','Use case','Precondition','Successful result'],[
['UC 01','Create workspace','No account exists','A valid account is created and signed in.'],
['UC 02','Import a capture','Authenticated user; valid bounded capture','A Ready session stores normalized replay events.'],
['UC 03','Replay traffic','A replayable session; no active worker','Packets and detections appear at the chosen speed.'],
['UC 04','Monitor live traffic','Administrator; interface and capture permission','Visible packets are analyzed until stopped or an error occurs.'],
['UC 05','Investigate an alert','An alert exists','Evidence is reviewed and status or notes are saved.'],
['UC 06','Tune detection','Valid thresholds and network settings','New sessions use the saved configuration.'],
['UC 07','Import EVE alerts','Valid newline delimited EVE JSON','An alert only session is available for triage.'],
['UC 08','Edit developer profile','Administrator; existing profile','Profile details and valid optional files are updated.'],
['UC 09','Export results','Authenticated user','CSV records or a JSON summary are downloaded.'],
['UC 10','Delete a session','The session is not active','Its retained packets, replay events, alerts, and notes are removed.']],[.6,1.35,2.1,2.5])
h('Replay use case detail')
p('The user imports a valid capture, selects a speed, and starts replay. The system freezes the active rule configuration, resets that session’s previous results, and schedules packet events. The browser displays processed counts and progress. The user can pause, change speed, resume, or stop. Completion changes the status to Finished. A processing error changes it to Error and records a readable explanation.')
h('Investigation alternative flow')
p('If supporting traffic is unavailable, the analyst can retain the New or Investigating status and document the uncertainty. An external EVE alert intentionally has no reconstructed packet evidence. A user can mark a reviewed finding Resolved or False positive without deleting its original trigger evidence.')

page('4 Software architecture')
p('The application is a layered monolith with an account scoped workspace resolver. Flask owns HTTP access, authentication, validation, rendering, and reporting. Each account engine owns at most one active replay or capture worker. Packet normalization and detection are separate modules so detector behavior can be tested without opening a capture interface.')
figure('02-components.png','Figure 2 Application components and persistence boundaries',max_h=5.8)
p('SQLite stores normalized metadata and alert evidence. Private local files hold the session secret, photos, and CVs. The EVE import path writes external alerts directly rather than running the behavioral detector. Browser updates use a snapshot stream and a packet polling endpoint. A separate read only integration supplies public GitHub repositories through a shared cache.')

page('5 Data flow and deployment')
figure('03-data-flow.png','Figure 3 Packet processing flow for live and replay sources',max_h=5.8)
p('A captured packet becomes a normalized event containing addresses, transport metadata, size, flags, and an optional decoded DNS name. Import sorting provides reproducible order. Detector windows advance using event timestamps. Each ingest transaction stores the packet, any alert evidence, and the processed count together.')
p('The live adapter buffers events in a bounded queue. An imported EVE alert follows a separate validation path and does not pass through packet normalization. Neither replay nor fixture generation transmits packets.')

page('5 Deployment environment')
figure('06-deployment.png','Figure 4 Default local deployment',max_h=4.5)
h('Runtime organization')
p('The launcher binds Waitress to 127.0.0.1 on port 5000 by default. Explicit HOST and TRUSTED_HOSTS configuration can enable trusted LAN access. It runs one application process with twenty four HTTP threads and account scoped engine workers. One active capture or replay session is allowed per account. More than one WSGI process must not share a data directory because engine coordination is held in process memory.')
h('Software and hardware assumptions')
p('Python 3.11 or newer is required. Release verification used Python 3.12 on Linux and Chromium for the browser checks. A typical student laptop is an appropriate starting environment for the bounded synthetic captures, but no minimum hardware throughput guarantee is claimed. RAM and disk use depend on imported event counts, retained sessions, and capture rate.')
h('Privilege boundary')
p('Offline parsing uses dpkt and requires ordinary file access. Scapy is loaded only when live capture is requested. Operating system capture permissions and Windows Npcap requirements belong to that optional path. A capture failure must not prevent the dashboard and replay features from starting.')

page('6 Account and workspace requirements')
req('FR 01','First run setup','The system shall present account creation when no user exists. It shall accept a Gmail address, a username of 3 to 40 characters, and a password of 12 to 128 characters. The first account shall be an administrator. It shall not ship a reusable default password.')
req('FR 02','Authentication','The system shall require authentication for dashboard pages, packet data, exports, private media, and mutation routes. An unauthenticated API request shall receive an authentication error rather than packet content.')
req('FR 03','Password storage','The system shall store a password hash using Werkzeug’s password hashing API. It shall not store a plaintext password in the user table or configuration.')
req('FR 04','Session behavior','A successful sign in shall replace the previous login session state. The session cookie shall be HttpOnly and SameSite Strict. The configured session lifetime shall be eight hours.')
req('FR 05','Form protection','Every POST request except static resource handling shall validate a token associated with the current session. An invalid or missing token shall cause a 400 response with a recovery message.')
req('FR 06','Login throttling','The application shall restrict login submissions to ten attempts per remote address in a rolling five minute interval. A blocked request shall return status 429. This throttle is local to the process.')
req('FR 07','Password update','An authenticated user shall be able to change the password after supplying the current password and a valid replacement. The update shall be recorded in the activity log.')
req('FR 08','Sign out','The sign out operation shall require a protected POST and clear the browser session. Subsequent protected requests shall require sign in.')
h('Setup race handling')
p('Account creation uses a local setup lock to avoid creating two initial administrators from concurrent setup requests in the same process. The supported one process deployment remains a condition of this behavior. Account recovery is deliberately not implemented as a hidden password or unauthenticated reset endpoint.')
h('Verification')
p('Web tests verify protected routes, setup behavior through fixtures, CSRF rejection, security headers, and password changes. Additional acceptance cases should include invalid credential attempts and checking that private CV downloads require a valid session.')

page('7 Capture ingestion requirements')
req('FR 09','Capture import','The system shall accept PCAP, PCAPNG, and CAP filenames subject to capture content validation. A file extension alone shall not establish that a file is valid.')
req('FR 10','Import limits','A capture shall be limited to 32 MiB and 50,000 capture records, including unsupported records counted toward the record limit. More than 100 stored sessions in one account shall require deleting an older session in that account before another import.')
req('FR 11','Supported decoding','The parser shall normalize supported IPv4 and IPv6 traffic from Ethernet, raw IP, loopback, and Linux cooked capture link types. Non IP and undecodable packet payloads shall not create invented packet records.')
req('FR 12','PCAPNG restrictions','The parser shall reject multi interface or multi section PCAPNG files and invalid block lengths. This prevents interpreting packets using another interface’s link type.')
req('FR 13','Normalized events','Each supported event shall include capture time, source and destination addresses, source and destination ports where available, protocol, captured length, TCP flags where available, and an optional decoded DNS question name.')
req('FR 14','Replay storage','Imported events shall be sorted stably by capture timestamp and stored with sequence numbers. A successful import shall create a Ready session and shall not automatically start monitoring.')
req('FR 15','Temporary source cleanup','The temporary uploaded source file shall be removed after parsing, including a rejected import. Normalized replay events shall remain until their session is deleted.')
req('FR 16','Guided fixture','The application shall offer a clearly labeled synthetic demo containing 501 packets. With the default configuration, completing it shall produce one alert from each of the six built in rules.')
h('Failure and recovery')
p('Empty, unsupported, over limit, or unreadable captures shall return an explanatory message and leave no partial imported session. User supplied names are bounded display labels; randomly generated temporary paths shall be used for the upload bytes. A 33 MiB HTTP request ceiling allows a 32 MiB file plus ordinary multipart overhead.')

page('8 Replay and session lifecycle')
figure('04-session-state.png','Figure 5 Replay lifecycle and interrupted session handling',max_h=3.5)
req('FR 17','Single active worker','The account engine shall reject a start request while another replay or live worker for that account is active, including a paused replay. Another account’s worker is scoped separately.')
req('FR 18','Timing controls','Replay shall support 0.25, 0.5, 1, 2, 4, 8, 16, and 32 times speed. Timing shall scale scheduling delays rather than event timestamps. Pause shall preserve scheduling progress; resume shall continue it.')
req('FR 19','Stop and restart','Stop shall retain completed observations. Restart shall clear the selected session’s previous packets, alerts, investigation notes, and progress before replaying its stored events. The UI shall request confirmation when prior results exist.')
req('FR 20','Rule snapshot','A session shall freeze rules and detection network settings when it starts. The snapshot shall be stored in the session record. Rule edits shall take effect on a new or restarted session.')
req('FR 21','Process restart','At application startup, previously Running or Paused sessions shall be marked Interrupted with an explanation. Detector state shall not be silently reconstructed from incomplete progress.')
req('FR 22','Live session control','Live capture shall require an administrator role and shall support start and stop. Pause and resume shall be rejected for live capture. A new live start shall create a new session rather than replaying captured records.')

page('9 Behavioral detection requirements')
h('Common event time semantics')
p('For an observation at time t and window W, the detector retains eligible observations with timestamps at least t minus W. Imported events are sorted. A live observation older than the detector’s current clock is stored as traffic but excluded from detection updates. An observation exactly at a threshold can trigger an alert.')
p('After firing, the same rule and detection key shall be suppressed for W event time seconds. Each session owns its detector state. State for inactive keys is pruned periodically. The six rules are deterministic threshold heuristics, not machine learned classifiers.')
req('FR 23','Port scan rule','For each source and destination pair, the detector shall count distinct destination ports observed in TCP or UDP packets during the configured window. The default threshold is 12 ports within 30 seconds, with High severity. Repeating one port shall not increase the distinct count.')
req('FR 24','Host scan rule','For each source, the detector shall count distinct destination IP addresses during the configured window. The default threshold is 10 hosts within 30 seconds, with High severity. The rule does not prove that a destination answered or was successfully scanned.')
req('FR 25','SYN burst rule','For each source, the detector shall count TCP packets whose SYN bit is set and ACK bit is not set. The default threshold is 80 packets within 10 seconds, with High severity. SYN ACK packets shall not contribute.')
h('Interpretation and false positives')
p('Legitimate service discovery, administration tools, monitoring, load tests, and connection retries can meet these thresholds. A SYN burst is not evidence of a successful flood or exhausted server resources because the release does not reconstruct all TCP handshakes or measure the target’s capacity. The alert name therefore uses the term indicator.')
h('Evidence contract')
p('A built in alert shall retain the rule identifier, severity, trigger timestamp, source and destination, observed value, threshold, window duration, and a readable explanation. The trigger destination is the packet that crossed the threshold; it does not necessarily describe every destination contributing to a source wide rule.')

page('10 DNS volume and integration requirements')
req('FR 26','DNS query burst','The detector shall count decoded UDP DNS queries by source. Decoded responses shall not contribute. The default threshold is 60 queries within 30 seconds, with Medium severity. DNS over HTTPS, DNS over TLS, and TCP DNS stream reconstruction are outside scope.')
req('FR 27','Outbound volume','The detector shall total captured frame lengths for a source inside a configured local CIDR range when the destination is outside all configured local ranges. The default threshold is 250,000 bytes within 60 seconds, with Medium severity. It is a fixed threshold rather than a learned normal baseline.')
req('FR 28','Destination indicators','The detector shall match destination IPs exactly against the configured indicator list. Its default threshold is one matching packet within 60 seconds, with Critical severity. The supplied documentation address is a demo indicator and shall not be described as a real malicious host.')
req('FR 29','Rule administration','The user shall be able to enable each rule and set a count or byte threshold from 1 to 100,000,000 and a window from 1 to 3,600 seconds. Invalid form values shall leave the saved rule configuration unchanged.')
req('FR 30','EVE import','The system shall import newline delimited Suricata EVE JSON alert events from a file up to 32 MiB and 50,000 lines. Each line shall be below 1 MB. Invalid alert records shall reject the import rather than create partial alerts.')
req('FR 31','External alert evidence','An EVE alert shall require a timezone aware timestamp, valid source and destination IPs, and signature details. The system shall preserve signature identity in evidence and map severities 1, 2, and 3 to High, Medium, and Low.')
req('FR 32','External session distinction','An EVE import shall create a Finished Suricata session containing alerts without packet reconstruction. Its alerts shall be triageable but shall not be replayed as packets.')
h('Configuration snapshots')
p('Local ranges and indicators are validated using the Python ipaddress module. The default local ranges cover common private IPv4 ranges; users can add valid IPv6 CIDRs. Saving configuration does not modify an active detector’s snapshot. The retained packet ceiling is an operational control that can be changed independently of frozen detector logic.')

page('11 Alert and host investigations')
figure('07-investigation.png','Figure 6 Investigation workflow and analyst disposition',max_h=4.7)
req('FR 33','Alert inbox','The system shall list alerts with detection name, severity, source, capture time, and status. It shall support text, severity, status, and session filters and paginate stored results.')
req('FR 34','Investigation context','An alert detail view shall display trigger evidence and up to 100 recent packets from the same source and session within the trigger window. Related signals shall be other alerts from that source in that session, limited to 30 displayed records.')
req('FR 35','Analyst disposition','The user shall be able to save New, Investigating, Resolved, or False positive status with notes up to 3,000 characters. The original trigger evidence shall remain available.')
req('FR 36','Host exploration','The host overview shall show up to 100 addresses ranked by sent plus received bytes. Host detail shall show up to 100 recent involving packets and up to 50 recent involving alerts. An IP observation shall not be asserted to identify a unique physical device.')

page('12 User interface requirements')
table(['Navigation page','Primary content'],[['Overview','Animated radar hero, stored traffic counters, capture time chart, protocol mix, recent detections, guided demo.'],['Live Traffic','Searchable and paginated packet table, counters, automatic first page updates, CSV export.'],['Threat Alerts','Filtered inbox and links to evidence based investigation.'],['PCAP Replay','Capture upload, live interface selection, EVE import, session progress and controls.'],['Host Explorer','Address cards, sent and received volume, and host investigation links.'],['Analytics','Capture time chart, protocol mix, rule breakdown, and top talkers.'],['Detection Rules','Six editable rule cards with enable switches, thresholds, and windows.'],['The Developers','Two editable biographies, skills, links, photos, and CV downloads.'],['GitHub Projects','Public developer repositories with filters, update information, and labeled cache states.'],['Settings','Own network settings, password, retention and activity; administrator account permissions.']],[1.25,5.3])
req('FR 37','Separate pages','Primary navigation items shall load separate routes and page documents. The application shall not implement all navigation as anchors into one long website page.')
req('FR 38','Live browser updates','The browser shall receive counters, charts, connection state, and session progress through SSE approximately once per second while connected. The unfiltered first traffic page shall refresh packet rows approximately every 2.5 seconds.')
req('FR 39','Display semantics','Historical packets shall keep their capture date. The header clock shall show current UTC separately. Charts shall be labeled as capture time distributions rather than current wall clock throughput.')
req('FR 40','Responsive motion','The layout shall adapt to desktop and phone widths, with a collapsible mobile sidebar and contained horizontally scrollable tables. Animation shall respect the browser’s reduced motion preference. Controls shall have readable labels and keyboard focus states.')
p('Alert inbox rows, host rankings, and rule breakdown lists are server rendered. Users can refresh or navigate to obtain updated lists; updating totals does not imply that every server rendered list is a continuously streamed table. The replay page refreshes its controls when a session status changes.')

page('13 Developer profiles and administration')
req('FR 41','Editable authors','The system shall provide two initial developer placeholders. Each shall support full name, contribution role, biography, comma separated skills, email, GitHub, LinkedIn, and portfolio fields. The supplied GitHub links shall be populated without inventing author biographies. It shall not invent personal identities or credentials.')
req('FR 42','Safe profile links','Nonempty web links shall require HTTP or HTTPS URLs with a network location. Displayed user text shall remain escaped. Profile edits shall be accessible only to an authenticated administrator.')
req('FR 43','Profile photo','An administrator shall be able to upload a photo up to 5 MiB and 20 megapixels. The server shall decode, bound dimensions to 1,000 by 1,000 pixels preserving aspect ratio, convert to RGB, and save a generated PNG filename.')
req('FR 44','CV upload','An administrator shall be able to upload a PDF CV up to 5 MiB. The server shall check the PDF envelope and serve accepted CVs as attachments through an authenticated media route. This check shall not be represented as malware scanning.')
req('FR 45','Profile persistence','Leaving a file field empty shall retain the previous file. A successful replacement shall remove the previous media file. A failed validation shall not partially overwrite the saved profile.')
req('FR 46','Environment settings','Settings shall accept 1 to 100 local CIDR ranges, up to 1,000 exact destination IP indicators, and a retained packet ceiling between 1,000 and 500,000. The default ceiling shall be 100,000.')
req('FR 47','Session retention','Deletion of a nonactive session shall remove its replay events, analyzed packets, alerts, and notes through foreign key cascades. An active session shall be stopped before deletion. Destructive deletion shall have a browser confirmation.')
req('FR 48','Activity history','The system shall record configuration changes, capture import/start events, alert updates, profile updates, password changes, deletion, and processing errors. Settings shall display the most recent 30 entries. The log is an activity record rather than a tamper proof forensic audit.')
req('FR 49','Exports','The system shall provide CSV downloads for packet and alert records and a JSON summary. Session filtering shall apply to exported observations. Spreadsheet formula prefixes in untrusted exported text shall be defused. CSV exports are limited to 100,000 newest records.')

page('14 Database design')
figure('05-database.png','Figure 7 SQLite entities and session relationships',max_h=3.0)
h('Persistence rules')
p('Each connection enables foreign keys. SQLite WAL mode supports concurrent readers while the worker writes. The implementation opens short lived connections per operation rather than sharing a connection across request and worker threads. Query arguments use placeholders for user controlled values.')
p('The sessions entity is the ownership root for analyzed packets, replay events, and alerts. Deleting a session cascades to those entities. Shared user identities and developer profiles are stored in the main database. Additional accounts use their own workspace databases for packets, sessions, alerts, config, and audit records. These are selected from the authenticated user ID, not a user supplied file path. The original first account retains its legacy workspace in the main database. Session deletion does not remove account identities or developer media.')
h('Transaction boundaries')
p('Capture import creates its session and replay records in one transaction. EVE import creates its session and alerts in one transaction. Packet ingest inserts one packet, any alerts produced from it, and the processed increment together. Profile files are handled separately from database storage, with validation and cleanup around the final update.')
h('Indexes')
p('The schema indexes packet session identifiers and source addresses and alert session identifiers. Primary and unique keys cover user identity, configuration keys, entity identifiers, and replay sequence identity. The current design is appropriate for a bounded local team service rather than a high volume historical data warehouse.')

page('14 Data dictionary')
table(['Entity','Key fields','Meaning and constraints'],[
['users','id; username; password; email; role','Shared identity key; unique username and canonical Gmail; password hash; admin or analyst. Legacy email may be null until linked.'],
['config','key; value','JSON documents named rules and settings. Validated through the application forms.'],
['sessions','id; name; source; status','Source is pcap, demo, live, or suricata. Status expresses Ready, Running, Paused, Finished, Stopped, Error, or Interrupted.'],
['sessions','total; processed; speed','Total supported replay events, processed event count, and selected scheduling multiplier. Live total is open ended and stored as zero.'],
['sessions','created; error; dropped; rule_snapshot','Creation time; readable worker failure; live events not processed at stop; frozen detector rules and environment.'],
['replay_packets','session_id; seq; event','Composite sequence key and serialized normalized event. Owned by a session.'],
['packets','id; session_id; ts','Packet row identity, owning session, original timestamp.'],
['packets','src; dst; sport; dport; proto','Address and transport metadata. Unavailable ports are zero. Protocol label is TCP, UDP, ICMP, or OTHER.'],
['packets','length; flags; dns','Captured byte length; decoded TCP flags; optional first DNS question name. No raw payload field.'],
['alerts','id; session_id; ts; rule; severity','Alert identity, owning session, trigger time, rule name, and assigned severity.'],
['alerts','src; dst; title; evidence','Trigger addresses, bounded title, and JSON containing observed value, threshold, window, and explanation.'],
['alerts','status; note','New, investigating, resolved, or false_positive; analyst notes up to 3,000 characters.'],
['profiles','id; name; role; bio; skills','Two initial profile identities and user editable professional information.'],
['profiles','email; github; linkedin; website; photo; cv','Validated contact links and generated media filenames, not arbitrary filesystem paths.'],
['audit','id; ts; action; detail','Workspace action record with timestamp and bounded detail text.']],[1.0,1.95,3.6])
h('Normalized event contract')
p('Replay event JSON additionally contains dns_response, a Boolean used to exclude DNS replies from query burst counts. Source timestamps are finite, nonnegative, and bounded to a supported calendar range. The dashboard formats event timestamps in UTC.')

page('15 External and internal interfaces')
table(['Method','Route','Purpose'],[
['GET POST','/setup and /login','First account creation and authentication.'],['POST','/logout','Clear browser session.'],
['GET','/ and named page routes','Render overview or a separate application page.'],
['GET','/api/snapshot','Current counters, protocols, capture buckets, recent alerts, and session metadata.'],
['GET','/api/stream','SSE JSON snapshots; stream reconnects periodically.'],
['GET','/api/packets','Up to 50 newest packets for the selected scope.'],
['POST','/replay/upload','Validate and import a bounded packet capture.'],
['POST','/replay/demo','Create a clearly labeled guided synthetic session.'],
['POST','/replay/eve','Validate and import existing Suricata alert records.'],
['POST','/sessions/{id}/{action}','Start, pause, resume, stop, or change replay speed.'],
['POST','/live/start','Open the selected capture interface and start a live session.'],
['GET POST','/alerts/{id}','Read evidence or save status and notes.'],
['GET','/hosts/{ip}','Read recent activity and alerts for an address.'],
['POST','/rules and /settings','Save validated rules or environment settings.'],
['GET POST','/team/{id}/edit','Edit a developer profile and optional files.'],
['GET','/media/{generated_name}','Authenticated photo display or CV download.'],
['GET','/export/{kind}','CSV packets, CSV alerts, or JSON summary.'],
['POST','/sessions/{id}/delete','Delete a nonactive session and its owned records.']],[.65,2.15,3.75])
h('Transport and validation')
p('Forms use standard URL encoded or multipart requests. Protected POSTs include csrf_token or the X CSRF Token header. Query parameter session selects a known integer session identifier. Invalid identifiers return a 400 response; unknown selected sessions return 404. The application does not expose arbitrary shell command execution or packet injection endpoints.')
p('An SSE message contains one JSON object prefixed by data and followed by an empty line. Streams last up to sixty snapshots before reconnecting, allowing authentication to be checked again. The default launcher is HTTP on localhost; public TLS termination is outside the baseline.')

page('15 Account and integration interfaces')
table(['Method','Route','Purpose'],[
['GET POST','/register','Create another account with a unique canonical Gmail.'],
['POST','/account/email','Link a Gmail once to an existing legacy account after password confirmation.'],
['POST','/accounts/{id}/role','Administrator changes another account to administrator or analyst.'],
['GET','/api/github/{login}','Read cached public repositories for one of the two configured developers.'],
['GET','/api/capture/status','Administrator refreshes capture backend and interface diagnostics.']],[.7,2.15,3.7])
h('Account scope and authorization')
p('All packet, alert, host, export, rule, session control, and snapshot routes select the signed in account’s database. A numeric record ID alone does not grant access to another account’s records. Public developer profiles and GitHub repository metadata are shared. Their mutation and live host capture have explicit administrator checks.')
h('Appearance and repository browser behavior')
p('The browser stores its chosen dark or light appearance locally and applies it before the main stylesheet paints. Theme changes redraw Canvas charts using theme colors. Repository filters operate on metadata returned by the local API, with text inserted as text rather than executable HTML.')
p('The GitHub API handler accepts only MeazbahMehedi and ItsMONIKAAKTER. The integration does not accept arbitrary remote hosts, upload captures, or modify repositories. It retains previous successful results during errors and honors a retry cooldown.')
h('Upgrade path')
p('Database migration adds Gmail and role fields and unique indexes without deleting legacy records. The original administrator can still sign in with the old username and password. Back up and copy the previous instance directory when moving into an extracted release folder; do not copy the virtual environment.')

page('16 Security and quality requirements')
req('NFR 01','Local access boundary','The default launcher shall bind only to localhost. Trusted Host validation shall accept the local hostnames configured by the application. Internet exposure is not a release acceptance condition.')
req('NFR 02','Untrusted content','Captured and imported strings shall be escaped in templates or inserted using DOM textContent. User controlled SQL values shall be bound parameters. File access shall use generated names resolved through known profile records.')
req('NFR 03','Browser policies','Responses shall include a same origin Content Security Policy, MIME sniffing protection, frame denial, same origin referrer policy, and no store caching. External script or font CDNs shall not be required.')
req('NFR 04','Bounded workload','Imports, queues, stored sessions, retained packet counts, and form fields shall have explicit limits. Reaching the retained packet ceiling shall stop processing with a readable error rather than silently deleting history.')
req('NFR 05','Reproducibility','Identical imported events and identical rule snapshots shall yield the same detection results regardless of supported replay speed. The bundled fixture coverage shall be testable without live network access.')
req('NFR 06','Usability','The ten page navigation shall remain understandable at desktop and phone sizes. Keyboard focus, labels, contrast, and reduced motion behavior shall support routine use. This is a design requirement rather than a formal accessibility certification.')
req('NFR 07','Maintainability','Packet parsing, detection, persistence, engine behavior, and web integration shall be organized into distinct Python modules. Meaningful tests shall cover their core contracts and regression boundaries.')
req('NFR 08','Observability','The interface shall distinguish local connection status, live monitoring, replay status, source type, progress, and failures. Queue drops shall not be represented as analyzed packets.')
req('NFR 09','Performance reporting','No unmeasured enterprise throughput or zero false positive guarantee shall be claimed. Performance acceptance for this release is successful processing of the bounded fixtures and responsive use in the verified environment.')
h('Privacy')
p('The runtime directory contains sensitive metadata and optional CVs. It shall be excluded from version control by default. The system shall not send packet data to external AI services. Raw payloads are not stored in packet records, but addresses and DNS names can still be sensitive.')

page('17 Operational constraints and failures')
table(['Condition','Required observable response'],[
['Invalid credentials','Generic sign in failure without revealing whether a username exists.'],
['Missing or stale CSRF token','400 response and instruction to refresh.'],
['Oversized HTTP request','413 response through the standard error page.'],
['Malformed or empty capture','Readable import failure and removal of the temporary file.'],
['Unsupported PCAPNG topology','Reject with a single interface export instruction.'],
['Another worker is active','Reject the new start request; retain the existing session.'],
['Live permission or interface failure','Explain that capture could not start; offline features remain available.'],
['Queue overflow or stop backlog','Count events that were not processed; persist the total at live session end.'],
['Retained packet ceiling','Set session Error with a message to delete old sessions or adjust the limit.'],
['Application restart during processing','Mark the previous session Interrupted and require an explicit restart.'],
['Browser connection interruption','Show reconnecting state; EventSource retries its stream.'],
['Attempt to delete an active session','Reject deletion until the session is stopped.']],[2.1,4.45])
h('Retention and backup')
p('The default runtime directory is instance under the launch directory. It contains the SQLite database, WAL state when active, signing key, private media, and temporary import directory. Back up this directory while the application is stopped. Deleting a session releases logical capacity, although SQLite may retain physical file space for reuse.')
h('Concurrency constraints')
p('Each workspace engine lock coordinates that account’s commands within one process. Each replay worker owns its detector and private store. HTTP readers use separate database connections. SSE streams consume HTTP threads; the default twenty four thread launcher is intended for a small local team. Large numbers of tabs, multiple WSGI workers, and high traffic enterprise monitoring require another architecture.')

page('18 Traceability and verification')
table(['Requirement group','Implementation','Verification evidence'],[
['FR 01 to 08','Flask protection, setup, login, password routes','test_app.py authentication, CSRF, headers, and password cases; manual sign in checks.'],
['FR 09 to 16','packets.py and engine import','test_packets.py malformed, normal, IPv6, PCAPNG, topology, sorting and fixture cases; upload web test.'],
['FR 17 to 22','Engine scheduler and live adapter','test_engine.py speed invariance, pause, resume, busy state, restart, interruption, frozen rules; mocked live permission failure.'],
['FR 23 to 29','Detector and configuration forms','test_detection.py thresholds, counters, expiry, suppression, disabled rules, DNS replies, SYN ACK and IPv6 ranges.'],
['FR 30 to 32','integrations.py and EVE upload','test_integrations.py valid import, bad records, alert rendering, and analytics compatibility.'],
['FR 33 to 36','Alert and host routes and templates','Web tests for alert details, triage updates and host detail; browser investigation inspection.'],
['FR 37 to 40','Jinja pages, CSS and JavaScript','Desktop Chromium checks of ten routes; mobile overflow and interaction checks; zero observed JavaScript exceptions.'],
['FR 41 to 49','Profile, settings, media, export and delete routes','Web tests for link validation, photo/CV handling, formula escaping, settings validation and cascade deletion.'],
['FR 50 to 59','Account resolver, theme script, capture module, GitHub client, launcher','Account uniqueness and isolation tests, role tests, migration tests, driver diagnostics, cache/error tests, and browser shutdown/theme checks.'],['NFR 01 to 09','Launcher, policies, limits, module boundaries','Source review, response header tests, speed invariance tests, fixture evaluation and browser checks.']],[1.15,2.1,3.3])
h('Verification method')
p('Automated tests check deterministic contracts and important failure paths. Browser verification checks rendered page behavior and layout. Source review establishes deployment boundaries and unsupported claims. No single method proves that the application has no remaining bugs; release acceptance combines these forms of evidence.')
h('Further environment checks')
p('A Windows installation and actual physical interface capture remain environment specific checks for the project team. Scapy 2.8.0 is included; Npcap must still be installed separately on Windows. The release test environment does not grant the required raw capture access. The supplied GitHub workflow is configured for four Python versions, but a successful run in the user’s repository must be observed after publishing.')

page('19 Release validation and demonstration')
table(['Check','Observed release result'],[['Automated Python suite','103 tests passed on Python 3.12 in the release environment.'],['Synthetic threat fixture','501 supported packets; six alerts covering all six built in rules.'],['Simple normal fixture','60 supported packets; no alerts with the default rules.'],['Desktop navigation','Ten primary pages returned successful responses and rendered in Chromium.'],['Browser JavaScript','No uncaught errors in both themes, ten pages, registration, triage, and profile workflows.'],['Responsive checks','Ten pages checked at desktop and 390 pixel phone sizes in both dark and light themes.'],['Physical live capture','Windows hardware capture not verified; driver, permission, interface and loopback paths tested.'],['Shutdown','Repeated Ctrl C during a paused replay and active SSE connection exited cleanly without a traceback.'],['Windows and macOS execution','Launch instructions and scripts supplied; native OS execution not performed.'],['Word document','Rendered and visually reviewed as part of release preparation.']],[2.0,4.55])
h('Reproducible commands')
p('From the project directory, install requirements dev, run python -m pytest -q, then run python scripts/evaluate.py. The fixture generator can recreate the bundled PCAP and EVE files without sending network traffic.')
h('Demonstration sequence')
p('Create the workspace account. Load the guided demo and start replay at a supported speed. Open the alert inbox after completion and verify six signals. Investigate the port scan, explain its distinct port threshold, save an analyst note, and view the source host. Show the analytics view, rules, external EVE import, developer profiles, and export controls.')
h('Meaning of the results')
p('The supplied fixtures are intentionally small and deterministic. Their results establish parser and detector behavior for those scenarios. They do not measure general precision, recall, resistance to evasive attacks, or production capacity. A broader labeled corpus and controlled workload measurements are needed for research claims.')

page('20 Installation and acceptance')
h('Installation procedure')
p('Extract the source ZIP. Install Python 3.11 or newer. Create a virtual environment, install requirements.txt with that environment’s Python, and run run.py. Open http://127.0.0.1:5000, create the first Gmail based administrator account, and complete the guided demo. Other users can register from the sign in page. Windows users can use start_windows.bat; Linux and macOS users can use start_unix.sh.')
p('requirements lock records the exact application dependency versions exercised for release verification. requirements dev adds pytest and optional browser tooling. The application does not need Node, React, a GPU, an API key, or an external database service.')
h('Acceptance checklist')
table(['ID','Acceptance condition'],[['AC 01','A fresh workspace creates an account and opens the dashboard without capture privileges.'],['AC 02','All ten navigation pages load, with distinct routes and meaningful empty states.'],['AC 03','The threat fixture completes at 1x, 4x, and 32x with the same six detector categories.'],['AC 04','Pause, resume, stop, restart, and interrupted status behave as specified.'],['AC 05','An alert exposes observed value, threshold, source, capture time, and context; a status/note edit persists.'],['AC 06','Invalid capture, configuration, profile link, and file inputs are rejected without partial saved updates.'],['AC 07','Profile photos and PDF CVs can be updated and retrieved only through authenticated routes.'],['AC 08','CSV and JSON exports are downloadable and respect selected observation scope.'],['AC 09','A nonactive session can be deleted with its dependent records, while an active session cannot.'],['AC 10','Tests and fixture evaluation complete, and native live capture is separately checked on the target demonstration machine if it will be presented.']],[.65,5.9])
h('Release handover')
p('The handover includes Python source, local frontend assets, synthetic fixtures, tests, GitHub CI configuration, a detailed README, security/contribution notes, this SRS, editable architecture diagram sources, and a presentation walkthrough. Personal authorship fields remain editable for the two developers.')

page('21 Limitations and future extensions')
h('Known practical limits')
p('The current parser stores packet metadata and supports selected link types; it is not a complete network protocol analyzer. DNS classification is limited to decoded UDP questions. Encrypted content is not decrypted. The host inventory tracks addresses rather than authenticated device identities. An alert suppression window can reduce repeated notifications but also hide repeated events within that interval.')
p('Rule thresholds require local tuning. A fixed outbound volume threshold cannot distinguish a legitimate backup from exfiltration on its own. The exact destination list is only as useful as its configured entries. A successful capture depends on interface visibility, application administrator access, OS permissions, and a functional capture driver. The release has no automated response or packet blocking capability. Gmail registration enforces uniqueness without verifying mailbox ownership.')
h('Prioritized future work')
p('The next useful extension is a labeled, licensed evaluation corpus covering normal and suspicious traffic from different environments. It should measure precision, recall, false positive rate, detection delay, and processing cost under documented conditions. This is a stronger foundation for an AI extension than attaching an unvalidated model.')
p('After establishing that baseline, the team can evaluate a small anomaly model against the rules, add richer flow reconstruction and Zeek logs, or separate sensors from the web server. Larger deployments would need authenticated sensor transport, multi process coordination, verified email and stronger role management, a scalable database, and measured retention policies.')
h('Change management')
p('Changes to thresholds, supported capture types, schema, state transitions, or external integrations should update requirements, fixtures, tests, README, and this document together. Requirements should retain stable identifiers so that a reviewer can trace behavior to code and acceptance evidence.')
h('Reference material')
for text in [
'Flask documentation. Security Considerations and Configuration. https://flask.palletsprojects.com/en/stable/',
'dpkt project documentation and source. https://dpkt.readthedocs.io/ and https://github.com/kbandla/dpkt',
'Scapy documentation. Asynchronous sniffing and API usage. https://scapy.readthedocs.io/',
'Suricata documentation. EVE JSON output. https://docs.suricata.io/en/latest/output/eve/eve-json-output.html',
'Python documentation. sqlite3, ipaddress, threading, and queue. https://docs.python.org/3/',
'SentinelFlow source code, test suite, README, validation record, and bundled synthetic fixtures, release 1.1.0.'
]:p(text)

page('22 Accounts appearance and upgrade requirements')
figure('08-account-isolation.png','Figure 8 Identity checks and account scoped storage',max_h=3.8)
req('FR 50','Gmail registration','The system shall allow additional accounts to register using a canonical Gmail address. Letter case, dots, plus tags, and the googlemail alias shall resolve to one mailbox key protected by a database unique index. No email ownership verification is included.')
req('FR 51','Private workspaces','Captured records, replay state, notes, rules, exports and settings shall be selected from the authenticated account’s store and engine. Another account’s IDs shall not expose or modify its records.')
req('FR 52','Account permissions','The first account shall be administrator and later accounts analysts. Only administrators shall capture host traffic, edit shared developer profiles, or change another account’s role. Self demotion shall be rejected.')
req('FR 53','Dark and light themes','Every page shall support dark and light appearance, with theme aware charts, first visit OS preference, saved browser selection, and a visible theme control. Theme choice shall persist across navigation.')
req('FR 58','Legacy migration','An existing first account and its captured data shall survive additive migration. A legacy user shall be able to link one Gmail by confirming the current password. Updated packages shall not require copying an old virtual environment.')

page('23 GitHub capture and shutdown requirements')
figure('09-github-cache.png','Figure 9 Public repository refresh and offline cache behavior',max_h=3.6)
req('FR 54','Developer repositories','A separate GitHub Projects page shall display public repositories for MeazbahMehedi and ItsMONIKAAKTER, with descriptions, languages, topics, stars, forks, update dates, and local filters. An account with no public projects shall have an honest empty state.')
req('FR 55','GitHub cache and errors','Successful public metadata shall be cached for five minutes. The page shall distinguish live refresh, cached, stale and unavailable states, enforce request bounds, and apply a cooldown after network or rate limit errors. Captures shall not be sent to GitHub.')
req('FR 56','Capture readiness','The live panel shall check the backend, show friendly interface labels, and explain missing Npcap or permissions. Windows shall load Npcap before initializing Scapy. A layer three sending socket shall not be presented as equivalent passive capture.')
req('FR 57','Cooperative shutdown','Ctrl C shall signal termination without raising inside Waitress shutdown waits. Engine workers and SSE streams shall be woken before HTTP cleanup. Repeated Ctrl C during cleanup shall not print a KeyboardInterrupt traceback.')
req('FR 59','Explicit LAN access','The default bind address shall remain localhost. Optional HOST and TRUSTED_HOSTS configuration shall be documented for a trusted LAN. This configuration shall not be described as a complete public internet deployment.')

out=ROOT/'docs'/'SentinelFlow_SRS.docx'
doc.save(out)
print(out)
